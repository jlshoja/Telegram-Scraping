import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

from tgcommon import extract_codes, extract_fields, normalize

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

POSTS_FILE = Path("cache/posts.json")
SRC = "Product Database.xlsx"
REPORT = "report.txt"

FIELDS = ["ابعاد", "جنس", "نوع بند", "کاربرد"]
COL = {  # 1-based excel columns
    "ابعاد": 11,
    "جنس": 12,
    "نوع بند": 13,
    "کاربرد": 14,
    "دسته بندی": 6,
}


def usage_to_category(usage):
    """دوشی | پاسپورتی  ->  کیف دوشی زنانه | کیف پاسپورتی زنانه"""
    parts = [p.strip() for p in str(usage).split("|") if p.strip()]
    out = []
    for p in parts:
        out.append(p if p.startswith("کیف") else f"کیف {p} زنانه")
    return " | ".join(dict.fromkeys(out))


def build_index(posts):
    exact = {}
    nums = {}
    for p in posts:
        fields = extract_fields(p["text"])
        ek, nk = extract_codes(p["text"])
        for key, mtype in ek.items():
            exact.setdefault(key, []).append((mtype, fields, p))
        for key, mtype in nk.items():
            nums.setdefault(key, []).append((mtype, fields, p))
    return exact, nums


def build_seller_affinity(rows, nums):
    """How often a seller's product codes show up in each channel."""
    counts = {}
    for _row, seller, num in rows:
        if not seller:
            continue
        for mtype, _fields, p in nums.get(num, ()):
            w = 3 if mtype == "code_kw" else 1
            counts[(seller, p["channel"])] = counts.get((seller, p["channel"]), 0) + w
    return counts


def score(mtype, fields):
    base = {"code_kw": 60, "full": 40}.get(mtype, 10)
    return base + 2 * len(fields)


def best_candidate(cands, affinity):
    def key(c):
        mtype, fields, post = c
        return (
            affinity.get(post["channel"], 0),
            score(mtype, fields),
            len(fields),
            len(post["text"]),
        )

    ranked = sorted(cands, key=key, reverse=True)
    top = key(ranked[0])
    ties = [c for c in ranked if key(c) == top]
    return ranked[0], len(ties)


def main(args):
    posts = json.loads(POSTS_FILE.read_text(encoding="utf-8"))
    print(f"loaded {len(posts)} posts")
    exact, nums = build_index(posts)

    wb = openpyxl.load_workbook(SRC)
    ws = wb.worksheets[0]
    header = {c.value: c.column for c in ws[1] if c.value}
    for name in FIELDS + ["کد محصول", "فروشنده", "دسته بندی"]:
        if name not in header:
            raise SystemExit(f"missing column: {name}")

    parsed = []
    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row, header["کد محصول"]).value
        seller = ws.cell(row, header["فروشنده"]).value
        code = normalize(str(raw)).strip().upper() if raw not in (None, "") else ""
        m = re.match(r"^([A-Z]+)-(\d+)$", code)
        parsed.append((row, code, m, seller))

    if args.codes:
        wanted = {c.strip().upper() for c in args.codes.split(",") if c.strip()}
        parsed = [p for p in parsed if p[1] in wanted]
        print(f"testing with codes: {sorted(wanted)}")

    code_rows = [(r, s, m.group(2)) for r, c, m, s in parsed if m]
    seller_affinity = build_seller_affinity(code_rows, nums)

    stats = {"filled_rows": 0, "no_code": 0, "invalid": [], "not_found": [],
             "found_empty": [], "kept_existing": [], "ambiguous": [], "filled": [],
             "category": []}

    for row, code, m, seller in parsed:
        if not code:
            stats["no_code"] += 1
            continue
        if not m:
            stats["invalid"].append((row, code))
            continue

        cands = exact.get(code) or nums.get(m.group(2))
        if not cands:
            hint = [
                f"{p['channel']} #{p['id']}"
                for p in posts
                if re.search(rf"(?<!\d){re.escape(m.group(2))}(?!\d)", normalize(p["text"]))
            ][:3]
            stats["not_found"].append((row, code, hint))
            continue

        affinity = {}
        if seller:
            for (sel, chan), w in seller_affinity.items():
                if sel == seller:
                    affinity[chan] = w

        (mtype, fields, post), ties = best_candidate(cands, affinity)
        if ties > 1:
            stats["ambiguous"].append((row, code, ties))

        if not fields:
            stats["found_empty"].append(
                (row, code, post["channel"], post["id"], post["text"][:200])
            )
            continue

        written = {}
        for name in FIELDS:
            if name not in fields:
                continue
            cell = ws.cell(row, COL[name])
            if cell.value not in (None, "") and not args.refresh:
                stats["kept_existing"].append((row, code, name, cell.value))
                continue
            cell.value = fields[name]
            written[name] = fields[name]
        if written:
            stats["filled_rows"] += 1
            stats["filled"].append((row, code, post["channel"], post["id"], written))

    for row in range(2, ws.max_row + 1):
        cat_cell = ws.cell(row, COL["دسته بندی"])
        if cat_cell.value not in (None, "") and not args.refresh:
            continue
        usage = ws.cell(row, COL["کاربرد"]).value
        if usage in (None, ""):
            continue
        cat_cell.value = usage_to_category(usage)
        stats["category"].append((row, usage, cat_cell.value))

    wb.save(SRC)
    print(f"saved {SRC}")

    lines = []
    lines.append(f"posts: {len(posts)}")
    lines.append(f"rows filled: {stats['filled_rows']}")
    lines.append(f"rows without code: {stats['no_code']}")
    lines.append("")
    lines.append("== seller -> channel used ==")
    for (sel, chan), w in sorted(seller_affinity.items(), key=lambda x: -x[1]):
        lines.append(f"  {sel} -> {chan} ({w})")
    lines.append("")
    lines.append(f"== filled ({len(stats['filled'])}) ==")
    for r, c, ch, mid, f in stats["filled"]:
        vals = "; ".join(f"{k}={v}" for k, v in f.items())
        lines.append(f"  row {r}: {c}  [{ch} #{mid}]  {vals}")
    lines.append(f"== دسته بندی derived from کاربرد ({len(stats['category'])}) ==")
    for r, u, c in stats["category"]:
        lines.append(f"  row {r}: {u}  ->  {c}")
    lines.append(f"== invalid code ({len(stats['invalid'])}) ==")
    for r, c in stats["invalid"]:
        lines.append(f"  row {r}: {c}")
    lines.append(f"== code not found in channels ({len(stats['not_found'])}) ==")
    for r, c, hint in stats["not_found"]:
        extra = f"  (digits seen in: {', '.join(hint)})" if hint else ""
        lines.append(f"  row {r}: {c}{extra}")
    lines.append(f"== found post but no field extracted ({len(stats['found_empty'])}) ==")
    for r, c, ch, mid, txt in stats["found_empty"]:
        lines.append(f"  row {r}: {c}  [{ch} #{mid}] {txt}")
    lines.append(f"== ambiguous candidates ({len(stats['ambiguous'])}) ==")
    for r, c, n in stats["ambiguous"]:
        lines.append(f"  row {r}: {c} ({n} tied)")
    lines.append(f"== existing values kept ({len(stats['kept_existing'])}) ==")
    for r, c, n, v in stats["kept_existing"]:
        lines.append(f"  row {r}: {c} {n}={v}")

    Path(REPORT).write_text("\n".join(lines), encoding="utf-8")
    print(f"report -> {REPORT}")
    print(f"filled={stats['filled_rows']} not_found={len(stats['not_found'])} "
          f"empty_post={len(stats['found_empty'])} invalid={len(stats['invalid'])}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true",
                    help="overwrite already-filled cells with fresh values")
    ap.add_argument("--codes", default="",
                    help="only process these product codes, comma-separated (for testing)")
    args = ap.parse_args()
    main(args)
